import re
from collections import Counter
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.hardware.engine import HardwareResult, resolve_hardware, spare_slot_hardware
from app.core.model.allocation import Allocation, Placement
from app.core.model.layout import Cioc
from app.core.model.pack import Pack
from app.core.model.signal import Signal
from app.core.naming.engine import Namer

DASH = "-"
STATUS = {
    "field": "Field",
    "wired_spare": "Wired Spare",
    "cabinet_alarm": "Cabinet Alarm",
    "spare_slot": "Spare Slot",
}


class WriterError(ValueError):
    pass


def _norm(h: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h).lower())


def _val(v: object) -> object:
    return DASH if v is None or v == "" else v


def dst_name(tag: str | None) -> str:
    return re.sub(r"[^A-Za-z0-9]", "", tag or "") + "_"


def unit_of(area: str | None, tag: str | None) -> object:
    for text in (tag or "", area or ""):
        m = re.search(r"\d{3,}", text)
        if m:
            return int(m.group(0))
    return DASH


def _unit_text(c: dict) -> str | None:
    u = unit_of(c["area"], c["tag"])
    return None if u == DASH else str(u)


def _majority(values: list) -> object:
    c = Counter(v for v in values if v not in (None, "", DASH))
    return min(c, key=lambda k: (-c[k], str(k))) if c else None


def _cabinet_context(alloc: Allocation) -> dict[int, dict]:
    """Technical room, operating area and risk area per cabinet, from the signals it holds."""
    ctx: dict[int, dict] = {}
    by_cab: dict[int, list[Signal]] = {}
    cab_of = {c.no: c.cabinet for c in alloc.layout.ciocs}
    for p in alloc.placements:
        if p.signal and p.kind == "field":
            by_cab.setdefault(cab_of[p.cioc_no], []).append(p.signal)
    every = [s for ss in by_cab.values() for s in ss]
    for cab in alloc.layout.cabinets:
        ss = by_cab.get(cab.no) or every
        ctx[cab.no] = {
            "technical_room": _majority([s.technical_room for s in ss]),
            "operating_area": _majority([s.operating_area for s in ss]),
            "area": _majority([s.area for s in ss]),
            "tag": _majority([s.tag for s in ss]),
        }
    return ctx


def _hardware(p: Placement | None, pack: Pack) -> HardwareResult:
    if p is None or p.kind == "spare_slot":
        return spare_slot_hardware(pack)
    probe = p.signal or Signal(
        io_type=p.io_type,
        duty_class=p.duty_class,
        signal_type=p.signal_type,
        signal_level=p.signal_level,
    )
    if p.kind == "cabinet_alarm" and p.signal:
        probe = p.signal
    return resolve_hardware(probe, pack)


def build_rows(
    alloc: Allocation, pack: Pack, batch: str = "B1", side: str = "F"
) -> list[dict[str, object]]:
    """One row per CHARM position, keyed by normalised header. The sheet is only a view."""
    cap = pack.capacity
    ch = cap.channels_per_baseplate
    layout = alloc.layout
    at = {(p.cioc_no, p.baseplate, p.channel): p for p in alloc.placements}
    ctx = _cabinet_context(alloc)
    ctl_of = {c.no: next(t for t in layout.controllers if c.no in t.ciocs) for c in layout.ciocs}
    cab_name = {c.no: c.name for c in layout.cabinets}
    namer = Namer(pack.naming)
    rows: list[dict[str, object]] = []
    for cab in layout.cabinets:
        for cno in cab.ciocs:
            cioc: Cioc = layout.cioc(cno)
            ctl = ctl_of[cno]
            for bp in cioc.baseplates:
                for channel in range(1, ch + 1):
                    p = at.get((cno, bp.n, channel))
                    s = p.signal if p else None
                    active = p if p and p.kind != "spare_slot" else None
                    hw = _hardware(p, pack)
                    c = ctx[cab.no]
                    status = STATUS[p.kind] if p else "Spare Slot"
                    internal = bool(s and s.internal_wiring)
                    terms = hw.tb_terminals + [DASH] * (3 - len(hw.tb_terminals))
                    tag = s.tag if s else DASH
                    if active and active.kind == "wired_spare" and pack.naming.wired_spare_tag:
                        tag = namer.name_for(
                            "wired_spare_tag",
                            unit=_unit_text(c),
                            seq=ctl.no,
                            cioc_no=cno,
                            n=bp.n,
                            ch=channel,
                        ).name
                    f = {
                        "srno": len(rows) + 1,
                        "revisionno": (s.revision if s and s.revision else "00")
                        if active and active.kind in ("field", "wired_spare")
                        else 0,
                        "tagname": tag,
                        "technicalroom": (s.technical_room if s else None) or c["technical_room"],
                        "operatingarea": (s.operating_area if s else None) or c["operating_area"],
                        "riskarea": (s.area if s else None) or c["area"],
                        "unit": unit_of(s.area, s.tag)
                        if s and s.area
                        else unit_of(c["area"], c["tag"]),
                        "segregation": s.segregation if s else None,
                        "dstname": dst_name(tag if status != "Spare Slot" else None),
                        "servicedescription": (s.service_short or s.service)
                        if s
                        else ("SPARE " + (active.io_type or "")[:2] if active else None),
                        "location": (s.extras.get("Location") if s else None),
                        "iotype": (s.io_type if s else (active.io_type if active else None)),
                        "signaltype": (
                            s.signal_type if s else (active.signal_type if active else None)
                        ),
                        "signallevel": (
                            s.signal_level if s else (active.signal_level if active else None)
                        ),
                        "nonisis": (s.execution if s else None)
                        or (None if not active else "NON-IS"),
                        "equipmentname": s.equipment_name if s else None,
                        "externalpower": s.external_power if s else None,
                        "systemcabinet": ctl.system_cabinet,
                        "controller": ctl.name,
                        "electronicmarshallingcabinet": cab_name[cab.no],
                        "cioccslsname": cioc.name,
                        "baseplate": bp.name,
                        "channelno": channel,
                        "charmno": (bp.n - 1) * ch + channel,
                        "charmtype": hw.charm,
                        "charmdescription": hw.charm_name,
                        "charmtbpartno": hw.tb,
                        "charmtbdescription": hw.tb_name,
                        "charmtb1": terms[0],
                        "charmtb2": terms[1],
                        "charmtb3": terms[2],
                        "fieldcableno": s.cable
                        if s and not internal
                        else (active.block_id if active and active.kind == "wired_spare" else None),
                        "fieldcablesize": s.cable_formation if s and not internal else None,
                        "fieldcablepair": (
                            s.pair_label
                            if s and not internal
                            else (
                                f"Pair {active.pair}"
                                if active and active.kind == "wired_spare"
                                else None
                            )
                        ),
                        "tagstatus": status,
                        "remarks": alloc.slot_remarks.get(f"{cno}/{bp.n}/{channel}")
                        or (
                            "Internal Cabinet Alarm."
                            if status == "Cabinet Alarm" and internal
                            else s.remarks
                            if s
                            else None
                        ),
                        "changetype": (active.change if active else None)
                        or ("Deleted" if f"{cno}/{bp.n}/{channel}" in alloc.slot_remarks else None)
                        or (s.change_type if s else None),
                        "batch": batch,
                        "hwlooptypical": hw.loop_typical if active else None,
                        "system2": s.extras.get("System 2") if s else None,
                        "oldtag": s.old_tag if s else None,
                        "instrumenttype": s.instrument_description if s else None,
                        "cabinetsidefr": side
                        if cioc.seq_in_cabinet <= cap.ciocs_per_cabinet // 2
                        else "R",
                        "cabinetnameforvisualization": f"{cab_name[cab.no]}({side if cioc.seq_in_cabinet <= cap.ciocs_per_cabinet // 2 else 'R'})",
                        "loopname": s.loop_name if s else None,
                        "sectionsubunit": None,  # not derivable from the MDB (manual in the sample)
                    }
                    rows.append(f)
    return rows


def write_allocation(
    alloc: Allocation,
    pack: Pack,
    template_path: str | Path,
    out_path: str | Path,
    extra_sheets: dict[str, list[list[object]]] | None = None,
    batch: str = "B1",
    side: str = "F",
) -> Path:
    """Write into a copy of the blank template. The template file is never touched."""
    template_path, out_path = Path(template_path), Path(out_path)
    if not template_path.is_file():
        raise WriterError(f"template not found: {template_path}")
    if template_path.resolve() == out_path.resolve():
        raise WriterError("the output path must differ from the template path")
    try:
        wb = load_workbook(template_path)
    except Exception as e:
        raise WriterError(f"cannot open template {template_path.name}") from e
    ws = wb.worksheets[0]
    headers = [c.value for c in ws[1]]
    cols = {_norm(h): i + 1 for i, h in enumerate(headers) if h}
    rows = build_rows(alloc, pack, batch, side)
    style_row = [copy(c._style) for c in ws[2]] if ws.max_row >= 2 else None
    for r, row in enumerate(rows, start=2):
        for key, col in cols.items():
            ws.cell(row=r, column=col, value=_val(row.get(key)) if key in row else DASH)
        if style_row and r > 237:
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col)._style = copy(style_row[col - 1])
    for r in range(len(rows) + 2, ws.max_row + 1):  # template rows we did not need
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).value = None
    last_row = max(2, len(rows) + 1)
    last_col = ws.cell(row=1, column=len(headers)).column_letter
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    ws.print_area = f"A1:{last_col}{last_row}"
    for title, data in (extra_sheets or {}).items():
        sh = wb.create_sheet(title[:31])
        for line in data:
            sh.append([_val(x) if x is not None else None for x in line])
        if data:
            sh.freeze_panes = "A2"
            for cell in sh[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for index, values in enumerate(zip(*data, strict=False), start=1):
                width = max(len(str(v or "")) for v in values)
                sh.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 60)
                for cell in sh[get_column_letter(index)]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
