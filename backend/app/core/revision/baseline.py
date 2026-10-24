"""Baselines for revisions: stored runs (SQLite) and allocation workbooks."""

import json
import re
import sqlite3
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.core.ingest.normaliser import parse_formation
from app.core.layout.reserve import spare_baseplates
from app.core.model.allocation import Allocation, Placement
from app.core.model.layout import Baseplate, Cabinet, Cioc, Controller, Layout
from app.core.model.pack import Pack
from app.core.model.signal import Signal

SCHEMA_VERSION = 1
SLOT = tuple[int, int, int]  # (cioc_no, baseplate, channel)
WORKBOOK_FIELDS = [
    "tag", "io_type", "signal_type", "signal_level", "cable", "pair", "cable_formation",
    "cable_pairs", "segregation",
]  # fmt: skip


class BaselineError(ValueError):
    pass


def skey(s: Signal) -> tuple:
    """Identity of a signal inside one run (the same tuple the validator uses)."""
    return (s.source_file, s.source_row, s.tag, s.cable, s.pair)


class Baseline(BaseModel):
    label: str = "00"
    run_id: str | None = None
    source: str = ""
    allocation: Allocation
    signals: list[Signal] = Field(default_factory=list)  # every signal of that run, placed or not
    # fields the baseline really carries; None means all (a stored run). A workbook has fewer.
    known_fields: list[str] | None = None

    @property
    def layout(self) -> Layout:
        return self.allocation.layout

    def slots(self) -> dict[tuple, Placement]:
        """Signal key -> the placement it holds (signals on the TQ list have none)."""
        return {
            skey(p.signal): p
            for p in self.allocation.placements
            if p.signal is not None and p.kind in ("field", "cabinet_alarm")
        }


# -- SQLite store ------------------------------------------------------------------------------
class RunStore:
    """Runs and their allocations. Creation is idempotent and stamped with a schema version."""

    def __init__(self, path: str = ":memory:"):
        self.db = sqlite3.connect(path)
        self.db.executescript(
            """
            CREATE TABLE IF NOT EXISTS schema_version (version INTEGER NOT NULL);
            CREATE TABLE IF NOT EXISTS runs (
                id TEXT PRIMARY KEY, label TEXT NOT NULL, parent_id TEXT,
                pack_hash TEXT, rules_hash TEXT, created_at TEXT, note TEXT
            );
            CREATE TABLE IF NOT EXISTS allocations (
                run_id TEXT PRIMARY KEY, layout_json TEXT NOT NULL,
                allocation_json TEXT NOT NULL, signals_json TEXT NOT NULL
            );
            """
        )
        if self.db.execute("SELECT COUNT(*) FROM schema_version").fetchone()[0] == 0:
            self.db.execute("INSERT INTO schema_version VALUES (?)", (SCHEMA_VERSION,))
        self.db.commit()

    def schema_version(self) -> int:
        return self.db.execute("SELECT MAX(version) FROM schema_version").fetchone()[0]

    def save_run(
        self,
        label: str,
        alloc: Allocation,
        signals: list[Signal],
        pack_hash: str = "",
        rules_hash: str = "",
        parent_id: str | None = None,
        note: str = "",
        created_at: str | None = None,
    ) -> str:
        if parent_id is not None and self._run(parent_id) is None:
            raise BaselineError(f"parent run '{parent_id}' does not exist")
        n = self.db.execute("SELECT COUNT(*) FROM runs").fetchone()[0] + 1
        run_id = f"run-{n:04d}"
        stamp = created_at or datetime.now(UTC).isoformat(timespec="seconds")
        with self.db:
            self.db.execute(
                "INSERT INTO runs VALUES (?,?,?,?,?,?,?)",
                (run_id, label, parent_id, pack_hash, rules_hash, stamp, note),
            )
            self.db.execute(
                "INSERT INTO allocations VALUES (?,?,?,?)",
                (
                    run_id,
                    alloc.layout.model_dump_json(),
                    alloc.model_dump_json(),
                    json.dumps([json.loads(s.model_dump_json()) for s in signals]),
                ),
            )
        return run_id

    def _run(self, run_id: str) -> tuple | None:
        return self.db.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone()

    def list_runs(self) -> list[dict]:
        cols = ("id", "label", "parent_id", "pack_hash", "rules_hash", "created_at", "note")
        rows = self.db.execute("SELECT * FROM runs ORDER BY id").fetchall()
        return [dict(zip(cols, r, strict=True)) for r in rows]

    def load_baseline(self, run_id: str) -> Baseline:
        run = self._run(run_id)
        if run is None:
            raise BaselineError(f"run '{run_id}' does not exist")
        _, layout_json, alloc_json, sig_json = self.db.execute(
            "SELECT * FROM allocations WHERE run_id=?", (run_id,)
        ).fetchone()
        alloc = Allocation.model_validate_json(alloc_json)
        alloc.layout = Layout.model_validate_json(layout_json)  # the stored layout is the layout
        return Baseline(
            label=run[1],
            run_id=run_id,
            source=f"run:{run_id}",
            allocation=alloc,
            signals=[Signal.model_validate(d) for d in json.loads(sig_json)],
        )


# -- allocation workbook -------------------------------------------------------------------------
def _norm(h: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h).lower())


def _text(v: object) -> str | None:
    if v is None:
        return None
    t = str(v).strip()
    return None if t in ("", "-") else t


def _int_at_end(text: str | None) -> int | None:
    m = re.search(r"(\d+)\s*$", text or "")
    return int(m.group(1)) if m else None


def baseline_from_workbook(path: str | Path, pack: Pack, label: str = "00") -> Baseline:
    """Rebuild a baseline from an allocation workbook (the template's first sheet).

    The layout is read back from the cabinet, controller, CIOC and baseplate columns, so a
    revision against a workbook reuses the layout that workbook was issued with.
    """
    cap = pack.capacity
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [_norm(h) for h in next(it, ())]
    need = ("tagstatus", "cioccslsname", "baseplate", "channelno", "electronicmarshallingcabinet")
    missing = [n for n in need if n not in head]
    if missing:
        raise BaselineError(f"not an allocation workbook, columns missing: {', '.join(missing)}")
    rows = [dict(zip(head, r, strict=False)) for r in it if any(v is not None for v in r)]
    wb.close()

    layout = Layout(spare_baseplates_per_cabinet=spare_baseplates(pack))
    cab_by_name: dict[str, Cabinet] = {}
    cioc_by_name: dict[str, Cioc] = {}
    ctl_by_name: dict[str, Controller] = {}
    for r in rows:
        cab, cio = _text(r["electronicmarshallingcabinet"]), _text(r["cioccslsname"])
        if not cab or not cio:
            raise BaselineError(f"row {r.get('srno')}: cabinet or CIOC name is empty")
        if cab not in cab_by_name:
            cab_by_name[cab] = Cabinet(no=len(cab_by_name) + 1, name=cab, ciocs=[])
            layout.cabinets.append(cab_by_name[cab])
        if cio not in cioc_by_name:
            cabinet = cab_by_name[cab]
            ctl_name = _text(r.get("controller")) or f"CTL-{cabinet.no}"
            if ctl_name not in ctl_by_name:
                ctl_by_name[ctl_name] = Controller(
                    no=len(ctl_by_name) + 1,
                    name=ctl_name,
                    system_cabinet=_text(r.get("systemcabinet")) or "",
                    ciocs=[],
                )
                layout.controllers.append(ctl_by_name[ctl_name])
            cioc = Cioc(
                no=len(cioc_by_name) + 1,
                name=cio,
                cabinet=cabinet.no,
                seq_in_cabinet=len(cabinet.ciocs) + 1,
                controller=ctl_by_name[ctl_name].no,
                baseplates=[],
            )
            cioc_by_name[cio] = cioc
            cabinet.ciocs.append(cioc.no)
            ctl_by_name[ctl_name].ciocs.append(cioc.no)
            layout.ciocs.append(cioc)
        cioc = cioc_by_name[cio]
        n = _int_at_end(_text(r["baseplate"]))
        if n is None:
            raise BaselineError(f"row {r.get('srno')}: baseplate '{r['baseplate']}' has no number")
        if all(b.n != n for b in cioc.baseplates):
            cioc.baseplates.append(Baseplate(name=str(r["baseplate"]), cioc_no=cioc.no, n=n))
    for c in layout.ciocs:
        c.baseplates.sort(key=lambda b: b.n)

    alloc = Allocation(layout=layout, solver="workbook")
    signals: list[Signal] = []
    used: set[tuple[int, int]] = set()
    for i, r in enumerate(rows, start=2):
        status = _text(r["tagstatus"])
        if status in (None, "Spare Slot"):
            continue
        cioc = cioc_by_name[_text(r["cioccslsname"])]
        n, ch = _int_at_end(_text(r["baseplate"])), int(r["channelno"])
        used.add((cioc.no, n))
        pair_label = _text(r.get("fieldcablepair"))
        pair = _int_at_end(pair_label)
        cable = _text(r.get("fieldcableno"))
        if status == "Wired Spare":
            alloc.placements.append(
                Placement(
                    cioc_no=cioc.no,
                    baseplate=n,
                    channel=ch,
                    kind="wired_spare",
                    block_id=cable,
                    pair=pair,
                    io_type=_text(r.get("iotype")),
                )
            )
            continue
        alarm = status == "Cabinet Alarm"
        seg = _text(r.get("segregation"))
        sig = Signal(
            tag=_text(r.get("tagname")), io_type=_text(r.get("iotype")),
            signal_type=_text(r.get("signaltype")), signal_level=_text(r.get("signallevel")),
            cable=None if alarm else cable, pair=None if alarm else pair,
            cable_formation=None if alarm else _text(r.get("fieldcablesize")),
            cable_pairs=None if alarm else parse_formation(r.get("fieldcablesize"))[0],
            pair_label=None if alarm else pair_label, segregation=seg,
            segregation_family=seg if seg and len(seg) == 1 else None,
            service=_text(r.get("servicedescription")), old_tag=_text(r.get("oldtag")),
            revision=str(r["revisionno"]) if r.get("revisionno") not in (None, "") else None,
            execution=_text(r.get("nonisis")), cabinet_alarm=alarm,
            internal_wiring=alarm and "internal" in (_text(r.get("remarks")) or "").lower(),
            loop_name=_text(r.get("loopname")), source_file=Path(path).name, source_row=i,
        )  # fmt: skip
        signals.append(sig)
        alloc.placements.append(
            Placement(
                cioc_no=cioc.no,
                baseplate=n,
                channel=ch,
                kind="cabinet_alarm" if alarm else "field",
                signal=sig,
                block_id=None if alarm else cable,
                pair=None if alarm else pair,
                io_type=sig.io_type,
            )
        )
    _mark_reserved(layout, alloc, used, spare_baseplates(pack), cap.alarm_baseplate_index)
    return Baseline(
        label=label,
        source=f"workbook:{Path(path).name}",
        allocation=alloc,
        signals=signals,
        known_fields=WORKBOOK_FIELDS,
    )


def _mark_reserved(layout: Layout, alloc: Allocation, used: set, spare: int, alarm_bp: int) -> None:
    alarm_cells = {(p.cioc_no, p.baseplate) for p in alloc.placements if p.kind == "cabinet_alarm"}
    for cab in layout.cabinets:
        first, last = layout.cioc(cab.ciocs[0]), layout.cioc(cab.ciocs[-1])
        for b in first.baseplates:
            if (first.no, b.n) in alarm_cells or b.n == alarm_bp:
                b.reserved = "alarm"
        tail = [b for b in last.baseplates if b.reserved is None][
            len(last.baseplates) - spare - 1 :
        ]
        tail = tail[-spare:] if spare else []
        for b in tail:
            if (last.no, b.n) not in used:
                b.reserved = "spare"
