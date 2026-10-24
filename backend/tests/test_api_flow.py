import os

import openpyxl
import pytest

from app.core.model.allocation import Allocation
from tests.api_helpers import (
    make_client,
    new_project,
    sample_template_or_skip,
    simple_rows,
    upload,
    xlsx_bytes,
)
from tests.test_tq_answers import HDR, cable_rows


def run_project(client, tmp_path, rows=None, **kw):
    pid = new_project(client, **kw)
    assert upload(client, pid, xlsx_bytes(tmp_path, rows or simple_rows())).status_code == 201
    r = client.post(f"/api/projects/{pid}/runs", json={})
    assert r.status_code == 201, r.text
    return pid, r.json()


def test_project_lifecycle(tmp_path):  # T9.1.1
    c = make_client(tmp_path)
    pid, run = run_project(c, tmp_path)
    assert run["status"] == "finished" and run["summary"]["allocated"] == 19
    assert run["summary"]["unaccounted"] == []
    rid = run["id"]
    a = c.get(f"/api/projects/{pid}/runs/{rid}/allocation")
    assert a.status_code == 200
    body = a.json()
    assert len(body["placements"]) >= 19 and body["rows"]
    # the response parses back into the core model
    Allocation.model_validate(
        {"layout": c.get(f"/api/projects/{pid}/runs/{rid}/layout").json()["layout"],
         "placements": body["placements"], "unplaced": body["unplaced"], "issues": body["issues"]}
    )  # fmt: skip
    p = c.get(f"/api/projects/{pid}").json()
    assert (p["uploads"], p["runs"]) == (1, 1)
    assert [r["id"] for r in c.get(f"/api/projects/{pid}/runs").json()] == [rid]


def test_all_read_endpoints(tmp_path):
    c = make_client(tmp_path)
    pid, run = run_project(c, tmp_path)
    base = f"/api/projects/{pid}/runs/{run['id']}"
    assert c.get(f"{base}/layout").json()["layout"]["cabinets"]
    assert c.get(f"{base}/groups").json()["groups"]
    rep = c.get(f"{base}/report").json()["report"]
    assert rep["kpi"]["allocated"] == 19 and all(r["passed"] for r in rep["rules"])
    imp = c.get(f"{base}/impact").json()
    assert imp["available"] is False and imp["reason"]
    assert c.get(f"{base}/tqs").status_code == 200


def test_download_by_run_id(tmp_path):  # T9.1.2
    sample_template_or_skip()
    c = make_client(tmp_path)
    pid, run = run_project(c, tmp_path)
    assert run["download_ready"] is True
    r = c.get(f"/api/projects/{pid}/runs/{run['id']}/download")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    out = tmp_path / "dl.xlsx"
    out.write_bytes(r.content)
    wb = openpyxl.load_workbook(out)
    assert "TQ list" in wb.sheetnames and "Run Summary" in wb.sheetnames


def test_download_without_template_is_a_clear_409(tmp_path, monkeypatch):
    monkeypatch.delenv("SUDOKU_DATA_DIR", raising=False)
    from fastapi.testclient import TestClient

    from app.config import Settings
    from app.main import create_app

    c = TestClient(create_app(Settings(db_path=str(tmp_path / "a.db"), cpsat_enabled=False)))
    pid, run = run_project(c, tmp_path)
    assert run["download_ready"] is False
    r = c.get(f"/api/projects/{pid}/runs/{run['id']}/download")
    assert r.status_code == 409 and r.json()["code"] == "output_unavailable"


def test_unknown_run_is_404_with_error_shape(tmp_path):  # T9.1.6
    c = make_client(tmp_path)
    pid = new_project(c)
    for tail in (
        "",
        "/allocation",
        "/layout",
        "/report",
        "/groups",
        "/impact",
        "/tqs",
        "/download",
    ):
        r = c.get(f"/api/projects/{pid}/runs/deadbeef{tail}")
        assert r.status_code == 404 and r.json()["code"] == "not_found"
        assert set(r.json()) == {"code", "message", "details"}


def test_cross_project_access_is_404(tmp_path):  # T9.1.7
    c = make_client(tmp_path)
    a, run = run_project(c, tmp_path)
    b = new_project(c, name="Other")
    for tail in ("", "/allocation", "/report", "/download"):
        assert c.get(f"/api/projects/{b}/runs/{run['id']}{tail}").status_code == 404
    assert c.get(f"/api/projects/{a}/runs/{run['id']}/allocation").status_code == 200


def test_run_needs_an_upload(tmp_path):
    c = make_client(tmp_path)
    pid = new_project(c)
    r = c.post(f"/api/projects/{pid}/runs", json={})
    assert r.status_code == 409 and r.json()["code"] == "no_uploads"


def test_mapping_values_rules_lint(tmp_path):
    c = make_client(tmp_path)
    rows = [[*r, f"X{r[2][-4:]}"] for r in cable_rows("320-CE-0001", 3, io="XX")]
    pid = new_project(c)
    upload(c, pid, xlsx_bytes(tmp_path, rows, header=[*HDR, "Cable Name"]), "mdb.xlsx")
    m = c.get(f"/api/projects/{pid}/mapping").json()
    assert m["sheets"][0]["matches"]["cable"]["column"] == "Cable No" and m["required"]
    bad = c.put(f"/api/projects/{pid}/mapping", json={"confirmed": {"mdb.xlsx": {"cable": "Nope"}}})
    assert bad.status_code == 422 and bad.json()["code"] == "invalid_mapping"
    ok = c.put(
        f"/api/projects/{pid}/mapping", json={"confirmed": {"mdb.xlsx": {"cable": "Cable Name"}}}
    )
    assert ok.json()["sheets"][0]["matches"]["cable"]["method"] == "confirmed"

    v = c.get(f"/api/projects/{pid}/values").json()
    assert [(u["field"], u["value"], u["count"]) for u in v["unknown"]] == [("io_type", "XX", 3)]
    assert "AIR" in v["allowed"]["io_type"]
    assert (
        c.put(f"/api/projects/{pid}/values", json={"io_types": {"XX": "NOPE"}}).status_code == 422
    )
    v2 = c.put(f"/api/projects/{pid}/values", json={"io_types": {"xx": "DIR"}}).json()
    assert v2["unknown"] == [] and v2["overrides"]["io_types"] == {"XX": "DIR"}
    run = c.post(f"/api/projects/{pid}/runs", json={}).json()
    assert run["summary"]["allocated"] == 3

    r = c.get(f"/api/projects/{pid}/rules").json()
    assert r["preset"] == "plant-default" and "strict" in r["presets"] and r["rules"]
    assert c.put(f"/api/projects/{pid}/rules", json={"preset": "nope"}).status_code == 422
    custom = [{"id": "SEP", "primitive": "separate",
               "parameters": {"a": "io_type = AIR", "b": "io_type = DIR"}}]  # fmt: skip
    r2 = c.put(f"/api/projects/{pid}/rules", json={"rules": custom}).json()
    assert r2["custom"] is True and [x["id"] for x in r2["rules"]] == ["SEP"]
    worse = c.put(f"/api/projects/{pid}/rules", json={"rules": [{"id": "X"}]})
    assert worse.status_code == 422 and worse.json()["code"] == "invalid_rules"
    lint = c.post(f"/api/projects/{pid}/lint", json={}).json()
    assert "findings" in lint and lint["dry_run"][0]["rule_id"] == "SEP"


def test_unimplemented_endpoints_say_so(tmp_path):
    c = make_client(tmp_path)
    pid = new_project(c)
    for tail in ("compile", "mine"):
        r = c.post(f"/api/projects/{pid}/rules/{tail}")
        assert r.status_code == 501 and r.json()["code"] == "not_implemented"


def test_answering_a_tq_changes_the_next_run(tmp_path):
    c = make_client(tmp_path)
    pid, run = run_project(
        c, tmp_path, cable_rows("320-CE-0001", 3, io="XX") + cable_rows("320-CE-0002", 4)
    )
    assert run["summary"]["held_out"] == 3
    tqs = c.get(f"/api/projects/{pid}/runs/{run['id']}/tqs").json()
    t = next(t for t in tqs if t["code"] == "TQ-VALUE")
    bad = c.post(f"/api/projects/{pid}/tqs/{t['id']}/answer", json={"option": "Do nothing"})
    assert bad.status_code == 422 and bad.json()["code"] == "invalid_answer"
    gone = c.post(f"/api/projects/{pid}/tqs/TQ-0000000000/answer", json={"option": "x"})
    assert gone.status_code == 404
    ok = c.post(
        f"/api/projects/{pid}/tqs/{t['id']}/answer",
        json={"option": "Map to a known value", "value": "DIR", "scope": "cable", "author": "ana"},
    )
    assert ok.status_code == 200 and ok.json()["tq"]["status"] == "answered"
    assert [h["author"] for h in ok.json()["history"]] == ["ana"]
    run2 = c.post(f"/api/projects/{pid}/runs", json={}).json()
    assert run2["summary"]["held_out"] == 0 and run2["summary"]["allocated"] == 7
    assert "ingest" in run2["summary"]["stages_reused"]
    by_status = c.get(f"/api/projects/{pid}/tqs", params={"status": "answered"}).json()
    assert [x["id"] for x in by_status] == [t["id"]]


@pytest.mark.skipif(not os.environ.get("SUDOKU_DATA_DIR"), reason="no sample data")
def test_sample_through_the_api(tmp_path, sample_mdb_path):
    c = make_client(tmp_path)
    pid = new_project(c, area="P30", unit="32001")
    r = upload(c, pid, sample_mdb_path.read_bytes(), sample_mdb_path.name)
    assert r.status_code == 201
    run = c.post(f"/api/projects/{pid}/runs", json={}).json()
    assert run["status"] == "finished" and run["summary"]["allocated"] == 189
    assert run["summary"]["unaccounted"] == []
