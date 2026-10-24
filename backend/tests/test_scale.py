"""Scale: the fast rule check agrees with the full one, and big plants finish validly in time."""

import json
import os
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch

import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

from app.config import Settings
from app.core.layout.blocks import build_blocks
from app.core.model.pack import Pack
from app.core.rules.schema import load_rules
from app.core.solve import cpsat, greedy
from app.core.solve.greedy import Placed, rule_rejections_full, solve_greedy
from app.core.synth import synthetic_signals
from app.core.writer.template import write_allocation
from tests.conftest import PACK_DIR
from tests.generators import mdb_specs
from tests.test_metamorphic import RULE_POOL, run

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import scale_test

PACK = Pack.load(PACK_DIR)

MORE_RULES = [
    {"id": "R-ASSIGN", "primitive": "assign", "scope": "io_type = AIR", "parameters": {"cioc": 1}},
    {"id": "R-RESERVE", "primitive": "reserve", "parameters": {"baseplate": 8, "cioc_seq": 1}},
    {"id": "R-QUOTA", "primitive": "quota",
     "parameters": {"per": "cioc", "max_percent": 60}, "scope": "io_type = DIR"},
    {"id": "R-ORDER", "primitive": "order", "scope": "io_type = DIR",
     "parameters": {"by": "logic_number"}},
    {"id": "R-SPAN", "primitive": "limit_span", "scope": "segregation_family = A",
     "parameters": {"unit": "cioc", "max": 2}},
    {"id": "R-BPCAP", "primitive": "capacity", "parameters": {"per": "baseplate", "max": 10}},
    {"id": "R-TIGHT", "primitive": "capacity", "parameters": {"per": "cioc", "max": 24}},
    {"id": "R-TIGHTCAB", "primitive": "capacity", "parameters": {"per": "cabinet", "max": 70}},
    {"id": "R-LOGIC2", "primitive": "co_locate", "scope": "logic_number in [L1, L2]",
     "parameters": {"by": "logic_number", "level": "cioc"}},
    {"id": "R-GROUP", "primitive": "co_locate", "parameters": {"by": "group", "level": "cabinet"}},
    {"id": "R-SEPSEG", "primitive": "separate", "parameters": {
        "a": "segregation_family = A", "b": "io_type = DIR", "level": "cioc"}},
    {"id": "R-SEPCAB", "primitive": "separate",
     "parameters": {"a": "io_type = AIR", "b": "io_type = DOR", "level": "cabinet"}},
]  # fmt: skip


@settings(max_examples=120)
@given(
    mdb_specs(awkward=False),
    st.lists(st.sampled_from([*RULE_POOL, *MORE_RULES]), max_size=5, unique_by=lambda r: r["id"]),
)
def test_fast_rule_check_gives_the_same_allocation_as_the_full_one(spec, defs):
    rules = load_rules(defs)
    signals = run(spec).allocatable
    blocks = build_blocks(signals, PACK)
    fast = solve_greedy(blocks, PACK, "AR1", "100", rules)

    def full(rs, ctx, placed, new, seen=None):
        return rule_rejections_full(
            rs, ctx, placed.all if isinstance(placed, Placed) else placed, new
        )

    with patch.object(greedy, "_rule_rejections", full):
        slow = solve_greedy(blocks, PACK, "AR1", "100", rules)
    assert fast.model_dump_json() == slow.model_dump_json()


def one(n: int) -> dict:
    cp = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "scale_test.py"), "--one", str(n)],
        check=True, capture_output=True, text=True, cwd=ROOT,
        env={**os.environ, "PYTHONPATH": str(ROOT)},
    )  # fmt: skip
    return json.loads(cp.stdout.strip().splitlines()[-1])


def test_1k_signals_within_target():  # T12.2.1
    r = one(1000)
    assert r["run_s"] < 15 and r["unplaced"] == 0 and r["violations"] == 0


def test_10k_signals_valid_within_limits_and_memory():  # T12.2.2 / T12.2.4 / T12.2.5
    r = one(10000)
    assert r["run_s"] < 60 and r["unplaced"] == 0 and r["violations"] == 0
    assert r["placed"] >= 10000
    assert r["peak_rss_mb"] < 1024
    assert r["events"][0] == "cpsat:skipped" and r["solver"] == "greedy"
    if "workbook_kb" in r:  # the workbook is written (and opens in test_writer's own checks)
        assert r["workbook_kb"] > 100


def test_a_tight_time_limit_falls_back_and_the_run_succeeds():  # T12.2.3
    rows = scale_test.forced_attempts()
    for row in rows:
        assert row["solver"] == "greedy" and row["unplaced"] == 0
        assert "cpsat:fallback" in row["events"] and "greedy:used" in row["events"]
        assert row["wall_s"] < 15


def test_oversized_model_is_declined_before_it_is_built():
    blocks = build_blocks(synthetic_signals(6000, seed=3), PACK)
    r = cpsat.solve_cpsat(blocks, PACK, "AR1", "100", None, hint=False)
    assert r.status == "SKIPPED" and "more than" in r.message
    assert r.wall_seconds < 5


def test_10k_workbook_opens(tmp_path):  # T12.2.5
    template = Settings.from_env().resolve_template()
    if template is None:
        pytest.skip("no allocation template available (SUDOKU_DATA_DIR)")
    r = run(scale_test.plant(10000))
    dest = tmp_path / "out.xlsx"
    write_allocation(r.allocation, r.pack, template, dest)
    wb = load_workbook(dest, read_only=True)
    assert wb.sheetnames and dest.stat().st_size > 100_000


def test_order_rule_fast_check_matches_the_full_one():
    """First-fit-decreasing puts the bigger cable first, so the order rule has to bite."""
    from tests.generators import CableSpec, MdbSpec

    spec = MdbSpec(
        (
            CableSpec("CBL-000", 20, 20, io="DI", logic="L2"),
            CableSpec("CBL-001", 8, 8, io="DI", logic="L1"),
            CableSpec("CBL-002", 6, 6, io="DI", logic="L3"),
        )
    )
    rules = load_rules([MORE_RULES[3]])
    blocks = build_blocks(run(spec).allocatable, PACK)
    fast = solve_greedy(blocks, PACK, "AR1", "100", rules)
    with patch.object(
        greedy,
        "_rule_rejections",
        lambda rs, c, p, n, s=None: rule_rejections_full(
            rs, c, p.all if isinstance(p, Placed) else p, n
        ),
    ):
        slow = solve_greedy(blocks, PACK, "AR1", "100", rules)
    assert fast.model_dump_json() == slow.model_dump_json()
    assert len({p.cioc_no for p in fast.placements if p.kind == "field"}) > 1  # the rule bit
