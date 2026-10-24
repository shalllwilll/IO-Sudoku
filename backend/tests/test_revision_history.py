import openpyxl
import pytest

from app.core.layout.blocks import build_blocks
from app.core.model.pack import Pack
from app.core.revision.baseline import RunStore
from app.core.revision.history import HistoryError, RevisionChain, commit_revision
from app.core.revision.impact import build_diff, build_impact, revision_sheets
from app.core.revision.solve import solve_revision
from app.core.solve.greedy import solve_greedy
from app.core.writer.template import write_allocation
from tests.test_greedy import cable

FAST = {"deterministic_time": 2.0}


@pytest.fixture
def pack(pack_dir):
    return Pack.load(pack_dir)


def edit(signals, pred, **kw):
    return [s.model_copy(update=kw) if pred(s) else s for s in signals]


def plant():
    return cable("C1", 8, 12) + cable("C2", 6, io="DIR") + cable("C3", 12) + cable("CA", 5, fam="A")


def make_chain(pack, path=":memory:"):
    """00 baseline; 01 adds a cable and re-types C2-3; 05A removes C3 and moves C2 (family B)."""
    store = RunStore(path)
    chain = RevisionChain(store)
    s0 = plant()
    alloc = solve_greedy(build_blocks(s0, pack), pack, "AR1", "100")
    run0 = store.save_run("00", alloc, s0, pack.hash)
    chain.start(run0, "00")
    b0 = store.load_baseline(run0)

    s1 = edit(s0, lambda x: x.tag == "C2-3", io_type="DOR") + cable("C9", 5)
    r1 = solve_solve(b0, s1, pack, "01")
    commit_revision(store, chain, b0, r1, s1, pack, "01")

    b1 = chain.baseline_of("01")
    s2 = [x for x in s1 if x.cable != "C3"]
    s2 = edit(s2, lambda x: x.cable == "C2", segregation_family="B")
    s2 = edit(s2, lambda x: x.tag == "C2-3", io_type="AIR")
    r2 = solve_solve(b1, s2, pack, "05A")
    commit_revision(store, chain, b1, r2, s2, pack, "05A")
    return store, chain, (s0, s1, s2), (r1, r2)


def solve_solve(base, signals, pack, label):
    return solve_revision(base, signals, pack, label=label, **FAST)


def test_compare_two_revisions_lists_the_differences(pack):
    _, chain, (s0, _s1, s2), _ = make_chain(pack)
    cmp = chain.compare("00", "05A")
    assert (cmp.from_label, cmp.to_label) == ("00", "05A")
    assert cmp.added == 5  # C9
    assert cmp.removed == 12  # C3
    kinds = {d.kind for d in cmp.deltas}
    assert {"added", "removed"} <= kinds
    assert cmp.added + cmp.removed + cmp.modified + cmp.moved + cmp.unchanged == len(
        {*(x.tag for x in s0), *(x.tag for x in s2)}
    )
    assert all(d.old_slot != "-" for d in cmp.deltas if d.kind == "removed")
    assert chain.compare("00", "00").unchanged == len(s0)


def test_signal_history_lists_two_changes_in_order(pack):
    _, chain, *_ = make_chain(pack)
    h = chain.signal_history("C2-3")
    assert [e.label for e in h] == ["01", "05A"]
    assert h[0].change == "Modified" and "io_type" in h[0].detail
    assert chain.signal_history("C9-1")[0].change == "Added"
    assert chain.signal_history("NO-SUCH-TAG") == []


def test_unknown_revision_is_a_clear_error(pack):
    _, chain, *_ = make_chain(pack)
    with pytest.raises(HistoryError, match=r"unknown revision '09'.*00, 01, 05A"):
        chain.compare("00", "09")
    with pytest.raises(HistoryError, match="unknown revision"):
        chain.baseline_of("77")


def test_chain_of_one_is_just_the_baseline(pack):
    store = RunStore()
    chain = RevisionChain(store)
    s = plant()
    run = store.save_run("00", solve_greedy(build_blocks(s, pack), pack, "AR1", "100"), s)
    chain.start(run, "00")
    summary = chain.summary()
    assert len(summary) == 1 and summary[0].label == "00" and summary[0].impact is None
    assert summary[0].parent_label is None
    cum = chain.cumulative()
    assert len(cum) == len(s) and all(r.changes == 0 and r.first_label == "00" for r in cum)
    assert chain.signal_history("C1-1") == []


def test_chain_summary_links_parents_and_carries_impact(pack):
    _, chain, *_ = make_chain(pack)
    s = chain.summary()
    assert [(e.label, e.parent_label) for e in s] == [("00", None), ("01", "00"), ("05A", "01")]
    assert s[1].impact.added == 5 and s[2].impact.deleted == 12
    assert [e.seq for e in s] == [1, 2, 3]


def test_impact_totals_add_up_for_every_revision(pack):
    _, chain, (_s0, s1, s2), _ = make_chain(pack)
    for entry, revised in zip(chain.summary()[1:], (s1, s2), strict=True):
        i = entry.impact
        assert i.categories == i.total_signals
        assert i.unchanged + i.modified + i.deleted == i.total_baseline
        assert i.unchanged + i.modified + i.added == i.total_revised == len(revised)


def test_impact_reports_cabinets_spares_and_move_reasons(pack):
    _store, chain, _, (_r1, r2) = make_chain(pack)
    i1, i2 = (e.impact for e in chain.summary()[1:])
    assert i1.added == 5 and i1.modified == 1 and i1.moved == 0 and i1.deleted == 0
    assert i1.spare_slots_after == i1.spare_slots_before - 5  # five new positions are taken
    assert i1.cabinets_used_before == i1.cabinets_used_after == 1
    assert i2.deleted == 12 and i2.moved == len(r2.moves)
    assert i2.moves_by_reason == {k: v for k, v in i2.moves_by_reason.items() if v}
    assert sum(i2.moves_by_reason.values()) == i2.moved
    assert i2.spare_slots_after > i2.spare_slots_before or i2.moved


def test_diff_has_one_row_per_changed_signal(pack):
    _store, chain, _, (_r1, r2) = make_chain(pack)
    b1 = chain.baseline_of("01")
    diff = build_diff(r2, b1)
    by = {}
    for d in diff:
        by.setdefault(d.change, []).append(d)
    assert len(by["Deleted"]) == 12 and all(
        d.new_slot == "-" and d.old_slot != "-" for d in by["Deleted"]
    )
    assert len({(d.tag, d.change) for d in diff}) == len(diff)  # one row each
    assert all(d.reason for d in diff)
    assert not any(d.change == "" for d in diff)
    assert {d.tag for d in diff} >= {x.signal.tag for x in r2.match.changes if x.kind == "modified"}


def test_sheets_match_the_json(pack, blank_template_path, tmp_path):
    _store, chain, _, (_r1, r2) = make_chain(pack)
    b1 = chain.baseline_of("01")
    impact, diff = build_impact(r2, b1, pack, "05A"), build_diff(r2, b1)
    sheets = revision_sheets(impact, diff)
    out = write_allocation(r2.allocation, pack, blank_template_path, tmp_path / "o.xlsx", sheets)
    wb = openpyxl.load_workbook(out)
    assert {"Revision Impact", "Revision Diff"} <= set(wb.sheetnames)
    got = {r[0]: r[1] for r in wb["Revision Impact"].iter_rows(min_row=2, values_only=True)}
    want = impact.model_dump()
    assert got["added"] == want["added"] and got["deleted"] == want["deleted"]
    assert got["moved"] == want["moved"] and got["total_signals"] == want["total_signals"]
    rows = list(wb["Revision Diff"].iter_rows(values_only=True))
    assert rows[0][:3] == ("Tag", "Change", "Cable") and len(rows) == len(diff) + 1
    assert [r[1] for r in rows[1:]] == [d.change for d in diff]


def test_cumulative_view_after_all_revisions(pack):
    _, chain, (_s0, _s1, s2), _ = make_chain(pack)
    cum = {r.tag: r for r in chain.cumulative()}
    assert set(cum) == {x.tag for x in s2}
    assert cum["C9-1"].first_label == "01" and cum["C9-1"].last_changed_label == "01"
    assert cum["C2-3"].changes == 2 and cum["C2-3"].last_changed_label == "05A"
    assert cum["C1-1"].first_label == "00" and cum["C1-1"].changes == 0
    assert all(r.slot != "-" for r in cum.values())


def test_records_survive_reopening_the_database(pack, tmp_path):
    path = str(tmp_path / "chain.db")
    make_chain(pack, path)
    chain = RevisionChain(RunStore(path))
    assert chain.labels() == ["00", "01", "05A"]
    assert chain.compare("01", "05A").removed == 12


def test_chain_guards(pack):
    store = RunStore()
    chain = RevisionChain(store)
    s = cable("C1", 3)
    alloc = solve_greedy(build_blocks(s, pack), pack, "AR1", "100")
    run = store.save_run("00", alloc, s)
    with pytest.raises(HistoryError, match="baseline first"):
        chain.add(run, "01", build_impact_stub(), [])
    with pytest.raises(Exception, match="does not exist"):
        chain.start("run-9999")
    chain.start(run, "00")
    with pytest.raises(HistoryError, match="already has a baseline"):
        chain.start(run, "00")
    with pytest.raises(HistoryError, match="already recorded"):
        chain.add(run, "00", build_impact_stub(), [])


def build_impact_stub():
    from app.core.revision.impact import Impact

    return Impact(label="x", baseline_label="00")
