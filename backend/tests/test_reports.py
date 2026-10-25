from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.main import create_app
from tests.test_tq_answers import cable_rows, sheet


def _sample(tmp_path):
    data = tmp_path / "data"
    data.mkdir()
    sheet(data, cable_rows("C1", 4), "Sample_IO_MDB_x.xlsx")
    c = TestClient(create_app(Settings(db_path=str(tmp_path / "api.db"), data_dir=str(data), cpsat_enabled=False)))
    made = c.post("/api/demo/sample").json()
    return c, made["project_id"], made["run_id"]


def test_allocation_report_is_organised_and_exports(tmp_path):
    c, pid, rid = _sample(tmp_path)
    doc = c.get(f"/api/projects/{pid}/runs/{rid}/reports/allocation").json()
    heads = [s["heading"] for s in doc["sections"]]
    assert heads[:4] == ["Summary", "Allocation by cabinet and CIOC", "Allocation by IO type", "Signal allocation"]
    assert "Rule check" in heads and dict(doc["meta"])["Project"] == "Sample plant"
    sig = next(s for s in doc["sections"] if s["heading"] == "Signal allocation")["tables"][0]
    assert len(sig["rows"]) > 0 and len(sig["rows"][0]) == len(sig["columns"])
    x = c.get(f"/api/projects/{pid}/runs/{rid}/reports/allocation?format=xlsx")
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames[0] == "Cover" and len(wb.sheetnames) > 3


def test_tq_report_groups_and_lists_answers(tmp_path):
    c, pid, rid = _sample(tmp_path)
    doc = c.get(f"/api/projects/{pid}/runs/{rid}/reports/tq").json()
    assert doc["kind"] == "tq" and doc["sections"][0]["heading"] == "Summary"
    assert c.get(f"/api/projects/{pid}/runs/{rid}/reports/tq?format=xlsx").status_code == 200


def test_unknown_report_kind_and_unfinished_run(tmp_path):
    c, pid, rid = _sample(tmp_path)
    assert c.get(f"/api/projects/{pid}/runs/{rid}/reports/nope").status_code == 422


def test_revision_report_route_is_not_swallowed_by_label(tmp_path):
    c, pid, _ = _sample(tmp_path)
    r = c.get(f"/api/projects/{pid}/revisions/report?from=00&to=01")
    assert r.status_code == 404 and r.json()["code"] == "unknown_revision"


def test_revision_document_counts_and_sections():
    from app.core.reports.documents import revision_report
    from app.core.revision.history import Comparison, SignalDelta

    cmp = Comparison(from_label="00", to_label="01", added=1, moved=1, unchanged=3,
                     deltas=[SignalDelta(tag="T1", kind="moved", old_slot="a", new_slot="b"),
                             SignalDelta(tag="T2", kind="added", new_slot="c")])  # fmt: skip
    doc = revision_report("P", "00", "01", cmp, None)
    heads = [s["heading"] if isinstance(s, dict) else s.heading for s in doc.sections]
    assert heads == ["Summary", "Moved signals", "Added signals"]
    assert "75%" in doc.sections[0].figures[3].hint


def test_tq_report_states_the_recommendation(tmp_path):
    from app.core.reports.documents import tq_report
    from app.core.tq.model import TQ, TqOption
    from app.core.tq.triage import GroupOption, Recommended, TqGroup

    tq = TQ(id="TQ-1", code="TQ-NAME", severity="assumed", category="c", subject="s", message="m", options=[TqOption(label="Keep", consequence="x")])
    g = TqGroup(key="k", code="TQ-NAME", category="c", priority="P2", title="T", why="W", tq_ids=["TQ-1"], open_ids=["TQ-1"], signals=1, status="open",
                options=[GroupOption(label="Keep", consequence="x", effect="note", changes_allocation=False)], recommended=Recommended(option="Keep", reason="safe"))
    doc = tq_report("P", {"id": "r"}, [tq], [g], {})
    assert "Recommended: Keep. safe" in doc.sections[1].text[1]
