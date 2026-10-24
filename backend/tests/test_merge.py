from collections import Counter

import pytest
from openpyxl import Workbook

from app.core.ingest.merge import merge_files
from app.core.ingest.reader import read_workbook
from app.core.model.pack import Pack

HDR = ["Tag No", "IO Type", "Cable No", "Pair No", "Segregation", "Status"]


@pytest.fixture
def pack(pack_dir):
    return Pack.load(pack_dir)


def write(path, rows, header=HDR):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def r(tag, io="AI", cable="320-CE-0001", pair=1, seg="-", status="New"):
    return [tag, io, cable, f"Pair {pair}", seg, status]


def codes(ss):
    return Counter(i.code for i in ss.issues)


def test_two_files_union(pack, tmp_path):
    a = write(tmp_path / "a.xlsx", [r("320-T-1", pair=1), r("320-T-2", pair=2)])
    b = write(tmp_path / "b.xlsx", [r("320-T-3", pair=3)])
    ss = merge_files([b, a], pack)
    assert [s.tag for s in ss.signals] == ["320-T-1", "320-T-2", "320-T-3"]
    assert ss.total_rows == ss.accounted == 3


def test_conflict_later_file_wins(pack, tmp_path):
    a = write(tmp_path / "a.xlsx", [r("320-T-1")])
    b = write(tmp_path / "b.xlsx", [r("320-T-9")])
    ss = merge_files([a, b], pack)
    assert [s.tag for s in ss.signals] == ["320-T-9"]
    assert ss.superseded[0].tag == "320-T-1"
    msg = next(i.message for i in ss.issues if i.code == "TQ-DATA")
    assert "a.xlsx" in msg and "b.xlsx" in msg
    assert ss.total_rows == ss.accounted == 2


def test_pair_beyond_capacity(pack, tmp_path):
    hdr = HDR + ["Cable Formation"]
    p = write(
        tmp_path / "a.xlsx",
        [r("T1", pair=30) + ["24PX1MM2/FR"], r("T2", pair=2) + ["24PX1MM2/FR"]],
        hdr,
    )
    ss = merge_files([p], pack)
    assert len(ss.signals) == 1 and len(ss.held_out) == 1
    assert "capacity" in ss.held_out[0].issues[0].message
    assert ss.total_rows == ss.accounted


def test_a_and_b_on_one_cable(pack, tmp_path):
    p = write(tmp_path / "a.xlsx", [r("T1", seg="A"), r("T2", pair=2, seg="B")])
    assert codes(merge_files([p], pack))["TQ-SEG"] == 1


def test_blank_io_type_held_out(pack, tmp_path):
    p = write(tmp_path / "a.xlsx", [r("T1", io=""), r("T2", pair=2)])
    ss = merge_files([p], pack)
    assert len(ss.held_out) == 1 and ss.held_out[0].issues[0].code == "TQ-DATA"


def test_duplicate_tag(pack, tmp_path):
    p = write(tmp_path / "a.xlsx", [r("T1"), r("T1", pair=2)])
    ss = merge_files([p], pack)
    assert len(ss.signals) == 2 and codes(ss)["TQ-DATA"] == 1


def test_deleted_does_not_clash_and_accounting(pack, tmp_path):
    p = write(tmp_path / "a.xlsx", [r("T1", status="DELETED"), r("T2"), r("T3", io="XYZ")])
    ss = merge_files([p], pack)
    assert len(ss.deleted) == 1 and len(ss.signals) == 1 and len(ss.held_out) == 1
    assert ss.total_rows == ss.accounted == 3


def test_sample_is_clean_and_accounted(pack, sample_mdb_path):
    ss = merge_files([sample_mdb_path], pack)
    assert not ss.blocking, [i.message for i in ss.issues if i.severity == "blocking"]
    assert ss.total_rows == ss.accounted == 190
    assert len(ss.deleted) == 1 and len(ss.signals) == 189 and not ss.held_out
    assert sum(s.internal_wiring for s in ss.signals) == 9
    print(Counter((i.code, i.severity) for i in ss.issues))


def test_sample_split_in_two_equals_original(pack, sample_mdb_path, tmp_path):
    (res,) = read_workbook(sample_mdb_path)
    cols, rows = res.summary.columns, [x.values for x in res.rows]

    def dump(name, part):
        wb = Workbook()
        ws = wb.active
        ws.append(cols)
        for v in part:
            ws.append([v.get(c) for c in cols])
        wb.save(tmp_path / name)
        return tmp_path / name

    a, b = dump("a.xlsx", rows[:100]), dump("b.xlsx", rows[100:])
    whole = merge_files([sample_mdb_path], pack)
    split = merge_files([a, b], pack)
    key = lambda ss: sorted((s.cable, s.pair, s.tag) for s in ss.signals)
    assert key(split) == key(whole)


def test_fault_injected_sample(pack, sample_mdb_path, tmp_path):
    (res,) = read_workbook(sample_mdb_path)
    cols, rows = res.summary.columns, [dict(x.values) for x in res.rows]
    rows[0]["I/O Type Name"] = None
    rows[1]["Pair Number"] = "Pair 40"
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    for v in rows:
        ws.append([v.get(c) for c in cols])
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    ss = merge_files([p], pack)
    assert len(ss.held_out) == 2
    assert ss.total_rows == ss.accounted == 190
