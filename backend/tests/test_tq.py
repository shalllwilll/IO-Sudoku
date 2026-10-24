import time

import openpyxl
import pytest
from pydantic import ValidationError

from app.core.ingest.merge import merge_files
from app.core.layout.blocks import build_blocks
from app.core.model.issue import Issue
from app.core.model.pack import Pack
from app.core.solve.greedy import solve_greedy
from app.core.tq.model import TQ, tq_from_issue, tq_id
from app.core.tq.store import TqStore, collect_tqs, tq_sheet, unaccounted
from app.core.writer.template import write_allocation


def issue(code="TQ-DATA", msg="bad", row=3):
    return Issue(code=code, severity="warning", message=msg, source_file="a.xlsx", source_row=row)


def test_create_and_read():
    s = TqStore()
    t = tq_from_issue(issue())
    assert s.add(t) is True
    got = s.get(t.id)
    assert got == t and got.severity == "assumed" and got.category == "missing data"


def test_stable_id_and_one_record():
    s = TqStore()
    assert tq_id("TQ-DATA", "a", 3) == tq_id("TQ-DATA", "a", 3)
    s.add(tq_from_issue(issue()))
    assert s.add(tq_from_issue(issue())) is False
    assert len(s.all()) == 1


def test_invalid_code():
    with pytest.raises(ValidationError):
        TQ(id="x", code="TQ-XYZ", severity="info", category="x", subject="s", message="m")


def test_columns_match_model():
    assert TqStore().columns() == list(TQ.model_fields)


def test_many_tqs_fast_and_grouped():
    s = TqStore()
    codes = ["TQ-DATA", "TQ-SEG", "TQ-CAP", "TQ-VALUE"]
    t0 = time.time()
    s.add_many([tq_from_issue(issue(codes[i % 4], f"m{i}", i)) for i in range(1000)])
    assert time.time() - t0 < 5 and len(s.all()) == 1000
    g = s.grouped()
    assert set(g) == {"missing data", "rule conflicts", "capacity issues"}
    assert len(g["missing data"]) == 500


@pytest.fixture
def pack(pack_dir):
    return Pack.load(pack_dir)


def test_invariant_on_fault_injected_sample(pack, sample_mdb_path, tmp_path):
    from openpyxl import Workbook

    from app.core.ingest.reader import read_workbook

    (res,) = read_workbook(sample_mdb_path)
    cols, rows = res.summary.columns, [dict(x.values) for x in res.rows]
    rows[0]["I/O Type Name"] = "XYZ"
    rows[1]["Pair Number"] = "Pair 40"
    rows[2]["Multipair Cable No."] = None
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    for v in rows:
        ws.append([v.get(c) for c in cols])
    p = tmp_path / "bad.xlsx"
    wb.save(p)

    ss = merge_files([p], pack)
    a = solve_greedy(build_blocks(ss.signals, pack), pack, "P30", "32001")
    held = [h.signal for h in ss.held_out if h.signal]
    tqs = collect_tqs(ss.issues, a, held)
    # held-out rows that never became Signals (unknown IO type) are named by their TQ too
    for h in ss.held_out:
        if h.signal is None:
            from app.core.tq.model import signal_key

            for t in tqs:
                if signal_key(h.source_file, h.source_row) in t.signal_keys:
                    break
            else:
                pytest.fail(f"held-out row {h.source_row} has no TQ")
    placed = sum(1 for x in a.placements if x.kind in ("field", "cabinet_alarm"))
    total = ss.total_rows - len(ss.deleted) - len(ss.out_of_scope) - len(ss.superseded)
    assert placed + len(a.unplaced) + len(ss.held_out) == total
    assert unaccounted([*ss.signals], a, tqs, held) == []


def test_tq_sheet_in_output(pack, blank_template_path, tmp_path):
    from tests.test_greedy import cable

    s = cable("C1", 3)
    a = solve_greedy(build_blocks(s, pack), pack)
    tqs = collect_tqs(a.issues, a)
    out = write_allocation(
        a, pack, blank_template_path, tmp_path / "o.xlsx", {"TQ list": tq_sheet(tqs)}
    )
    ws = openpyxl.load_workbook(out)["TQ list"]
    assert ws.max_row == len(tqs) + 1 and len(tqs) >= 1
