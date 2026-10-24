"""Damaged inputs: the run must finish, say what is wrong, and never lose a signal."""

import tempfile
from dataclasses import replace
from pathlib import Path

import pytest
from hypothesis import given
from hypothesis import strategies as st
from openpyxl import Workbook, load_workbook

from app.core.pipeline import PipelineRun
from app.core.reports import build_report
from app.core.solve import cpsat
from tests.generators import (
    HEADERS,
    CableSpec,
    MdbSpec,
    columns,
    load_fixture,
    mdb_specs,
    write_mdb,
)
from tests.test_metamorphic import PACK, cells, check_invariant, run

PLAIN = mdb_specs(awkward=False)
BASIC = MdbSpec(
    (
        CableSpec("CBL-000", 12, 12, io="DI", seg="A"),
        CableSpec("CBL-001", 6, 8, io="DO", duty="SOV"),
        CableSpec("CBL-002", 4, 4, io="AI", seg="B", logic="L1"),
    )
)


def pipeline(paths) -> PipelineRun:
    from app.core.pipeline import Pipeline

    return Pipeline(paths, PACK, "AR1", "100", cpsat_enabled=False).run()


def report_of(r: PipelineRun):
    issues = [*r.signalset.issues, *r.blocks.issues, *r.allocation.issues]
    return build_report(
        r.allocation,
        r.pack,
        r.signalset.signals,
        issues,
        r._held_signals(),
        r.not_enforced,
        r.rules,
    )


def assert_survived(r: PipelineRun):
    assert r.unaccounted() == []
    assert cpsat.gate(r.allocation, r.blocks, r.pack, r.rules) == []
    assert report_of(r).kpi is not None  # the report is produced whatever was wrong


def write(spec, folder) -> list[Path]:
    return write_mdb(spec, Path(folder))


@pytest.mark.parametrize("missing", ["tag", "io", "cable", "pair"])
def test_a_missing_required_column_is_a_blocking_tq_map(tmp_path, missing):  # T12.1.3
    paths = write_mdb(replace(BASIC, drop=(missing,)), tmp_path)
    r = pipeline(paths)
    maps = [t for t in r.tqs if t.code == "TQ-MAP" and t.severity == "blocking"]
    assert maps and any(field_name(missing) in t.message for t in maps)
    assert_survived(r)
    # nothing silently disappeared: the signals that could not be built are on the TQ list
    assert r.allocation.placements == [] or missing == "tag"


def field_name(missing: str) -> str:
    return {"io": "io_type"}.get(missing, missing)


@given(
    PLAIN,
    st.lists(
        st.tuples(
            st.integers(0, 400),
            st.integers(0, 12),
            st.one_of(
                st.none(),
                st.text(st.characters(blacklist_categories=("Cc", "Cs")), max_size=40),
                st.integers(-(10**6), 10**6),
                st.floats(allow_nan=False, allow_infinity=False),
                st.sampled_from(["=1+1", "#REF!", "  ", "-", "N/A", "☃", "A" * 300]),
            ),
        ),
        max_size=25,
    ),
)
def test_garbage_in_cells_gives_tqs_not_a_crash(spec, edits):  # T12.1.4
    with tempfile.TemporaryDirectory() as d:
        paths = write(spec, d)
        for path in paths:
            wb = load_workbook(path)
            ws = wb.active
            for row, col, value in edits:
                ws.cell(
                    row=2 + row % max(1, ws.max_row), column=1 + col % ws.max_column
                ).value = value
            wb.save(path)
        r = pipeline(paths)
        assert_survived(r)


@given(PLAIN, st.lists(st.integers(0, 300), min_size=1, max_size=15))
def test_duplicate_rows_never_double_allocate(spec, picks):
    with tempfile.TemporaryDirectory() as d:
        (path,) = write(replace(spec, files=1), d)
        wb = load_workbook(path)
        ws = wb.active
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        for i in picks:
            ws.append(rows[i % len(rows)])
        wb.save(path)
        r = pipeline([path])
        assert_survived(r)
        placed = [p.signal.tag for p in r.allocation.placements if p.kind == "field"]
        assert len(placed) == len(set(placed))  # a tag is placed at most once


@given(PLAIN, st.integers(1, 6))
def test_blank_rows_change_nothing(spec, blanks):
    assert cells(run(replace(spec, blank_rows=0)).allocation) == cells(
        run(replace(spec, blank_rows=blanks)).allocation
    )


def test_empty_workbook_gives_an_empty_valid_result(tmp_path):  # T12.1.5
    wb = Workbook()
    wb.save(tmp_path / "empty.xlsx")
    r = pipeline([tmp_path / "empty.xlsx"])
    assert r.allocation.placements == [] and r.allocation.unplaced == []
    assert_survived(r)


def test_header_only_workbook_gives_an_empty_valid_result(tmp_path):  # T12.1.5
    wb = Workbook()
    wb.active.append([HEADERS["default"][c] for c in columns(BASIC)])
    wb.save(tmp_path / "header.xlsx")
    r = pipeline([tmp_path / "header.xlsx"])
    assert r.allocation.placements == [] and r.allocation.unplaced == []
    assert_survived(r)


def test_no_files_at_all_is_an_empty_valid_result():  # T12.1.5
    r = pipeline([])
    assert r.allocation.placements == [] and r.allocation.unplaced == []
    assert_survived(r)


@pytest.mark.parametrize("io", ["DI", "DO", "AI", "AO"])
def test_a_single_signal_is_valid(tmp_path, io):  # T12.1.6
    spec = MdbSpec((CableSpec("ONE", 1, 1, io=io),))
    r = pipeline(write(spec, tmp_path))
    check_invariant(spec, r)
    assert any(p.kind == "field" for p in r.allocation.placements)
    assert_survived(r)


@pytest.mark.parametrize("name", ["alt_split_composite", "alt_combined_multifile", "alt_sparse"])
def test_alternate_schema_fixtures_allocate_like_the_plain_form(tmp_path, name):  # T12.1.2
    spec = load_fixture(name)
    r = pipeline(write(spec, tmp_path / "alt"))
    check_invariant(spec, r)
    assert r.allocation.unplaced == []
    plain = replace(
        spec,
        header_style="default",
        composite_io=False,
        combined_grouping=False,
        drop=(),
        extra_columns=0,
        files=1,
        blank_rows=0,
        shuffle_seed=None,
    )
    twin = pipeline(write(plain, tmp_path / "plain"))
    placed = lambda a: sorted((p.signal.tag, p.cioc_no) for p in a.placements if p.kind == "field")
    # composite IO types carry a duty the plain form lacks, so compare where the form is the same
    if not spec.composite_io:
        assert placed(r.allocation) == placed(twin.allocation)


# regressions found by the generators ---------------------------------------------------------
def test_pair_zero_is_held_out_with_a_tq_not_dropped(tmp_path):
    from tests.test_tq_answers import cable_rows, sheet

    rows = cable_rows("C-0001", 3)
    rows[0][3] = 0  # Pair No
    r = pipeline([sheet(tmp_path, rows)])
    assert_survived(r)
    held = [t for t in r.tqs if t.code == "TQ-DATA" and "pair number 0" in t.message]
    assert held and len(held[0].signal_keys) == 1
    assert sum(1 for p in r.allocation.placements if p.kind == "field") == 2


def test_two_signals_on_one_pair_do_not_lose_one(tmp_path):
    from tests.test_tq_answers import cable_rows, sheet

    rows = cable_rows("C-0001", 3)
    rows[2][3] = "Pair 2"
    r = pipeline([sheet(tmp_path, rows)])
    assert_survived(r)
    assert any("pair 2 appears in" in t.message for t in r.tqs)
    assert sum(1 for p in r.allocation.placements if p.kind == "field") == 2
