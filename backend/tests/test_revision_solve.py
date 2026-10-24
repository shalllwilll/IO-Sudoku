import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from app.core.ingest.merge import merge_files
from app.core.layout.blocks import build_blocks
from app.core.model.pack import Pack
from app.core.revision.baseline import RunStore, baseline_from_workbook
from app.core.revision.solve import REASONS, solve_revision
from app.core.rules.schema import load_rules
from app.core.solve import cpsat
from app.core.solve.greedy import solve_greedy
from app.core.solve.validator import validate_allocation
from app.core.writer.template import write_allocation
from tests.conftest import PACK_DIR
from tests.test_greedy import cable

FAST = {"deterministic_time": 2.0}


@pytest.fixture
def pack(pack_dir):
    return Pack.load(pack_dir)


def baseline_of(signals, pack, rules=None):
    alloc = solve_greedy(build_blocks(signals, pack), pack, "AR1", "100", rules)
    store = RunStore()
    return store.load_baseline(store.save_run("00", alloc, signals, pack.hash)), alloc


def plant():
    return cable("C1", 8, 12) + cable("C2", 6, io="DIR") + cable("C3", 12) + cable("CA", 5, fam="A")


def cells_of(alloc):
    return {
        (p.signal.tag, p.cioc_no, p.baseplate, p.channel)
        for p in alloc.placements
        if p.signal is not None
    }


def revise(base, revised, pack, rules=None, **kw):
    return solve_revision(base, revised, pack, rules, **{**FAST, **kw})


def edit(signals, pred, **kw):
    return [s.model_copy(update=kw) if pred(s) else s for s in signals]


def valid(res, pack, signals):
    assert res.violations == []
    # the invariant: every signal is allocated exactly once or is on the TQ list
    placed = [p.signal.tag for p in res.allocation.placements if p.signal]
    listed = [u.signal.tag for u in res.allocation.unplaced]
    assert sorted(placed + listed) == sorted(s.tag for s in signals)


def test_no_change_means_zero_moves(pack):
    s = plant()
    base, alloc = baseline_of(s, pack)
    res = revise(base, [x.model_copy() for x in s], pack)
    assert res.moves == [] and res.solver == "cpsat" and res.tier == 1
    assert cells_of(res.allocation) == cells_of(alloc)
    assert res.pinned_blocks == len(build_blocks(s, pack).blocks) and res.free_blocks == 0
    assert {p.kind for p in res.allocation.placements if p.block_id == "C1"} == {
        "field", "wired_spare",
    }  # fmt: skip
    valid(res, pack, s)


def test_added_cable_places_new_signals_and_moves_nothing(pack):
    s = plant()
    base, alloc = baseline_of(s, pack)
    new = cable("C9", 5)
    res = revise(base, [*s, *new], pack)
    assert res.moves == []
    old = cells_of(alloc)
    assert old <= cells_of(res.allocation)  # nothing existing moved
    added = {p.signal.tag for p in res.allocation.placements if p.change == "Added"}
    assert added == {x.tag for x in new}
    valid(res, pack, [*s, *new])


def test_new_signals_fill_the_spare_pairs_of_an_existing_cable(pack):
    s = plant()  # C1 uses 8 of 12 pairs
    base, alloc = baseline_of(s, pack)
    fill = cable("C1", 12, 12)[8:]
    res = revise(base, [*s, *fill], pack)
    assert res.moves == []
    spare = {
        (p.cioc_no, p.baseplate, p.channel)
        for p in alloc.placements
        if p.block_id == "C1" and p.kind == "wired_spare"
    }
    took = {
        (p.cioc_no, p.baseplate, p.channel) for p in res.allocation.placements if p.signal in fill
    }
    assert took == spare and all(
        p.change == "Added" for p in res.allocation.placements if p.signal in fill
    )
    valid(res, pack, [*s, *fill])


def test_every_move_has_a_reason_from_the_list(pack):
    s = plant()
    base, _ = baseline_of(s, pack)
    # C2 turns out to carry family B next to family A; C3 grows to 20 pairs
    rev = edit(s, lambda x: x.cable == "CA", segregation_family="A")
    rev = edit(rev, lambda x: x.cable == "C2", segregation_family="B")
    rev = [x for x in rev if x.cable != "C3"] + cable("C3", 20, 24)
    res = revise(base, rev, pack)
    assert res.moves, "this revision forces moves"
    assert all(m.reason in REASONS and m.detail for m in res.moves)
    assert {m.block_id for m in res.moves} <= {"C2", "C3"}  # stable blocks stay put
    reasons = {m.block_id: m.reason for m in res.moves}
    assert reasons.get("C3") == "block_growth"
    print(res.move_report())
    valid(res, pack, rev)


def test_modified_signal_with_placement_impact_moves_only_its_block(pack):
    base_sigs = cable("CA", 12, fam="A") + cable("CN", 12) + cable("CZ", 12)
    base, alloc = baseline_of(base_sigs, pack)
    same_cioc = {p.cioc_no for p in alloc.placements if p.block_id in ("CA", "CN")}
    rev = edit(base_sigs, lambda x: x.cable == "CN", segregation_family="B")
    res = revise(base, rev, pack)
    valid(res, pack, rev)
    moved = {m.block_id for m in res.moves}
    assert moved <= {"CN"}
    if len(same_cioc) == 1:  # A and B cannot share the CIOC they used to share
        assert moved == {"CN"} and all(m.reason == "modified" for m in res.moves)


def test_hardware_only_change_keeps_the_slot(pack):
    s = plant()
    base, alloc = baseline_of(s, pack)
    rev = edit(s, lambda x: x.tag == "C2-3", io_type="DOR")
    res = revise(base, rev, pack)
    assert res.moves == [] and cells_of(res.allocation) == cells_of(alloc)
    p = next(p for p in res.allocation.placements if p.signal and p.signal.tag == "C2-3")
    assert p.change == "Modified"


def test_no_room_goes_to_tq_cap_without_a_crash(pack):
    full = [x for i in range(28) for x in cable(f"F{i:02d}", 12)]
    base, alloc = baseline_of(full, pack)
    assert not alloc.unplaced
    rev = [*full, *cable("EXTRA", 6)]
    res = revise(base, rev, pack)
    assert {u.signal.cable for u in res.allocation.unplaced} == {"EXTRA"}
    assert all(i.code == "TQ-CAP" for u in res.allocation.unplaced for i in u.issues)
    assert res.moves == []
    valid(res, pack, rev)


def test_cpsat_failure_falls_back_to_greedy_pin(pack, monkeypatch):
    def boom(*a, **k):
        raise cpsat.CpsatError("broken")

    monkeypatch.setattr(cpsat, "_staged_solve", boom)
    s = plant()
    base, alloc = baseline_of(s, pack)
    new = cable("C9", 7)
    res = revise(base, [*s, *new], pack)
    assert res.solver == "greedy-pin" and any(e.status == "fallback" for e in res.events)
    assert cells_of(alloc) <= cells_of(res.allocation) and res.moves == []
    valid(res, pack, [*s, *new])


def test_greedy_pin_alone_when_cpsat_is_off(pack):
    s = plant()
    base, alloc = baseline_of(s, pack)
    res = revise(base, [*s, *cable("C9", 4)], pack, use_cpsat=False)
    assert (
        res.solver == "greedy-pin"
        and res.moves == []
        and cells_of(alloc) <= cells_of(res.allocation)
    )


def test_deleting_cables_leaves_spare_slots_with_remarks(pack):
    s = plant()
    base, alloc = baseline_of(s, pack)
    gone = [x for x in s if x.cable == "C2"]
    res = revise(base, [x for x in s if x.cable != "C2"], pack, label="05A")
    freed = {(p.cioc_no, p.baseplate, p.channel) for p in alloc.placements if p.block_id == "C2"}
    assert len(res.allocation.slot_remarks) == len(freed) == 6 + 0  # six signals, no spare pairs
    assert all(
        r.startswith("Deleted in revision 05A") for r in res.allocation.slot_remarks.values()
    )
    assert any(g.tag in r for g in gone for r in res.allocation.slot_remarks.values())
    assert res.moves == []
    valid(res, pack, [x for x in s if x.cable != "C2"])


def test_deleted_rows_become_spare_slot_rows_in_the_workbook(pack, blank_template_path, tmp_path):
    import openpyxl

    s = [x.model_copy(update={"pair_label": f"Pair {x.pair}"}) for x in plant()]
    base, _ = baseline_of(s, pack)
    rev = [x for x in s if x.cable != "C2"]
    res = revise(base, rev, pack, label="02")
    out = write_allocation(res.allocation, pack, blank_template_path, tmp_path / "r.xlsx")
    ws = openpyxl.load_workbook(out).worksheets[0]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, [c.value for c in r], strict=False)) for r in ws.iter_rows(min_row=2)]
    deleted = [
        r for r in rows if r["Tag Status"] == "Spare Slot" and "Deleted" in str(r["Remarks"])
    ]
    assert len(deleted) == 6 and all("Deleted in revision 02" in r["Remarks"] for r in deleted)
    assert all(r["Change Type"] == "Deleted" for r in deleted)


def test_renamed_cable_matches_by_tag_keeps_the_slots_and_raises_tq_rev(pack):
    s = plant()
    base, alloc = baseline_of(s, pack)
    rev = edit(s, lambda x: x.cable == "C3", cable="C3R")
    res = revise(base, rev, pack)
    assert sum(i.code == "TQ-REV" for i in res.allocation.issues) == 12
    old = {c for c in cells_of(alloc) if c[0].startswith("C3-")}
    new = {c for c in cells_of(res.allocation) if c[0].startswith("C3-")}
    assert old == new and res.moves == []  # default: keep the slot
    assert {
        p.block_id for p in res.allocation.placements if p.signal and p.signal.tag == "C3-1"
    } == {"C3R"}
    valid(res, pack, rev)


def test_rule_change_reasons_are_reported(pack):
    s = cable("C1", 12) + cable("C2", 12)
    base, alloc = baseline_of(s, pack)
    first = next(p for p in alloc.placements if p.block_id == "C1")
    rules = load_rules(
        [{"id": "RV", "primitive": "assign", "scope": "cable = C1",
          "parameters": {"cioc": first.cioc_no, "baseplates": [first.baseplate + 3]}}]
    )  # fmt: skip
    res = revise(base, [x.model_copy() for x in s], pack, rules, rules_changed=True)
    assert res.moves and {m.block_id for m in res.moves} == {"C1"}
    assert all(m.reason == "rule_change" for m in res.moves)
    assert validate_allocation(res.allocation, pack, s, rules=rules) == []


def test_revision_is_deterministic(pack):
    s = plant()
    base, _ = baseline_of(s, pack)
    rev = [*s, *cable("C9", 9)]
    a, b = revise(base, rev, pack), revise(base, rev, pack)
    assert a.allocation.model_dump_json() == b.allocation.model_dump_json()


@settings(max_examples=12, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(
    st.lists(
        st.tuples(
            st.sampled_from(["add", "delete", "retype", "family"]),
            st.integers(0, 3),
            st.integers(1, 9),
        ),
        min_size=1,
        max_size=3,
    )
)
def test_stability_unchanged_blocks_never_move(edits):
    pack = Pack.load(PACK_DIR)
    s = cable("B0", 9) + cable("B1", 12) + cable("B2", 5, 8) + cable("B3", 7, fam="A")
    base, _ = baseline_of(s, pack)
    rev, touched, extra = list(s), set(), 0
    for kind, which, n in edits:
        cab = f"B{which}"
        if kind == "add":
            extra += 1
            rev += cable(f"N{extra}", n)
        elif kind == "delete":
            rev = [x for x in rev if x.cable != cab]
            touched.add(cab)
        elif kind == "retype":
            rev = edit(rev, lambda x, c=cab: x.cable == c, io_type="DIR")
        else:
            rev = edit(rev, lambda x, c=cab: x.cable == c, segregation_family="B")
            touched.add(cab)
    res = revise(base, rev, pack, deterministic_time=1.0)
    assert res.violations == []
    assert {m.block_id for m in res.moves} <= touched  # only blocks that really changed move
    assert all(m.reason in REASONS for m in res.moves)


def test_revising_the_issued_sample_against_itself(pack, sample_alloc_path, sample_mdb_path):
    base = baseline_from_workbook(sample_alloc_path, pack)
    mdb = merge_files([sample_mdb_path], pack).signals
    res = revise(base, mdb, pack)
    print(res.match.counts(), len(res.moves), res.solver, res.tier)
    print(res.move_report())
    assert res.violations == [] and res.solver in ("cpsat", "greedy-pin")
    assert res.pinned_blocks + res.free_blocks == len(build_blocks(mdb, pack).blocks)
    placed = sum(1 for p in res.allocation.placements if p.kind in ("field", "cabinet_alarm"))
    assert placed + len(res.allocation.unplaced) == sum(
        1 for c in res.match.changes if c.new is not None and c.kind != "deleted"
    )
