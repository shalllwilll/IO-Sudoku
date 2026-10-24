"""Revisions over REST: baseline, a revision run, the chain, comparison and signal history."""

import pytest

from tests.api_helpers import (
    make_client,
    new_project,
    run_now,
    sample_template_or_skip,
    simple_rows,
    upload,
    xlsx_bytes,
)
from tests.test_tq_answers import cable_rows


@pytest.fixture
def client(tmp_path):
    return make_client(tmp_path)


def put_mdb(client, tmp_path, pid, rows, name="mdb.xlsx"):
    """Replace the project's uploads with one workbook."""
    for u in client.get(f"/api/projects/{pid}/uploads").json():
        assert client.delete(f"/api/projects/{pid}/uploads/{u['id']}").status_code == 204
    r = upload(client, pid, xlsx_bytes(tmp_path, rows, name), name)
    assert r.status_code == 201, r.text


def baselined(client, tmp_path):
    pid = new_project(client)
    put_mdb(client, tmp_path, pid, simple_rows())
    run = run_now(client, pid)
    assert run["status"] == "finished"
    r = client.post(f"/api/projects/{pid}/baseline", json={"run_id": run["id"]})
    assert r.status_code == 201, r.text
    return pid, run


def revise(client, tmp_path, pid, rows, **body):
    put_mdb(client, tmp_path, pid, rows, name="rev.xlsx")
    r = client.post(f"/api/projects/{pid}/revisions", json=body)
    assert r.status_code == 201, r.text
    done = client.get(f"/api/projects/{pid}/runs/{r.json()['id']}").json()
    assert done["status"] == "finished", done
    return done


def revised_rows():
    rows = simple_rows() + cable_rows("320-CE-0003", 6)
    rows[2][4] = "A"  # a signal of the first cable changes its segregation
    return rows


def test_revision_needs_a_baseline(client, tmp_path):
    pid = new_project(client)
    put_mdb(client, tmp_path, pid, simple_rows())
    r = client.post(f"/api/projects/{pid}/revisions", json={})
    assert r.status_code == 409 and r.json()["code"] == "no_baseline"
    assert client.get(f"/api/projects/{pid}/revisions").json() == []


def test_baseline_heads_the_chain_once(client, tmp_path):
    pid, run = baselined(client, tmp_path)
    chain = client.get(f"/api/projects/{pid}/revisions").json()
    assert [e["label"] for e in chain] == ["00"] and chain[0]["impact"] is None
    again = client.post(f"/api/projects/{pid}/baseline", json={"run_id": run["id"]})
    assert again.status_code == 409 and again.json()["code"] == "baseline_exists"
    ghost = client.post(f"/api/projects/{pid}/baseline", json={"run_id": "nope"})
    assert ghost.status_code == 404 and ghost.json()["code"] == "not_found"


def test_revision_run_adds_to_the_chain_with_impact(client, tmp_path):  # T13.1.6
    pid, base = baselined(client, tmp_path)
    done = revise(client, tmp_path, pid, revised_rows())
    assert done["summary"]["revision"] == "01"
    imp = client.get(f"/api/projects/{pid}/runs/{done['id']}/impact").json()
    assert imp["available"] is True
    i = imp["impact"]
    assert i["label"] == "01" and i["baseline_label"] == "00"
    assert i["added"] == 6 and i["modified"] >= 1 and i["deleted"] == 0
    assert i["added"] + i["modified"] + i["deleted"] + i["unchanged"] == i["total_signals"]
    assert "added" in {d["change"].lower() for d in imp["diff"]}

    chain = client.get(f"/api/projects/{pid}/revisions").json()
    assert [e["label"] for e in chain] == ["00", "01"]
    assert chain[1]["parent_label"] == "00" and chain[1]["impact"]["added"] == 6
    one = client.get(f"/api/projects/{pid}/revisions/01").json()
    assert one["run_id"] == chain[1]["run_id"]
    assert client.get(f"/api/projects/{pid}/revisions/1").json()["label"] == "01"  # normalised

    alloc = client.get(f"/api/projects/{pid}/runs/{done['id']}/allocation").json()
    assert alloc["unplaced"] == [] and alloc["placements"]
    # the initial run still says it has no revision impact
    first = client.get(f"/api/projects/{pid}/runs/{base['id']}/impact").json()
    assert first["available"] is False and first["impact"] is None


def test_compare_matches_the_impact_and_unknown_labels_are_404(client, tmp_path):
    pid, _ = baselined(client, tmp_path)
    revise(client, tmp_path, pid, revised_rows())
    cmp = client.get(f"/api/projects/{pid}/revisions/compare", params={"from": "00", "to": "01"})
    assert cmp.status_code == 200, cmp.text
    c = cmp.json()
    assert c["from_label"] == "00" and c["to_label"] == "01" and c["added"] == 6
    assert len(c["deltas"]) == c["added"] + c["removed"] + c["modified"] + c["moved"]
    bad = client.get(f"/api/projects/{pid}/revisions/compare", params={"from": "00", "to": "09"})
    assert bad.status_code == 404 and bad.json()["code"] == "unknown_revision"
    gone = client.get(f"/api/projects/{pid}/revisions/09")
    assert gone.status_code == 404 and "recorded" in gone.json()["message"]
    missing = client.get(f"/api/projects/{pid}/revisions/compare", params={"from": "00"})
    assert missing.status_code == 422 and missing.json()["code"]


def test_signal_history_follows_a_tag_across_revisions(client, tmp_path):
    pid, _ = baselined(client, tmp_path)
    revise(client, tmp_path, pid, revised_rows())
    changed = simple_rows()[2][0]
    added = "320-DI-000301"
    hist = client.get(f"/api/projects/{pid}/signals/history", params={"tag": changed}).json()
    assert [h["label"] for h in hist] == ["01"] and hist[0]["change"].lower() in (
        "modified",
        "moved",
    )
    new = client.get(f"/api/projects/{pid}/signals/history", params={"tag": added}).json()
    assert [h["change"].lower() for h in new] == ["added"]
    assert client.get(f"/api/projects/{pid}/signals/history", params={"tag": "NOPE"}).json() == []
    assert client.get(f"/api/projects/{pid}/signals/history").status_code == 422


def test_cumulative_lists_every_current_signal(client, tmp_path):
    pid, _ = baselined(client, tmp_path)
    revise(client, tmp_path, pid, revised_rows())
    rows = client.get(f"/api/projects/{pid}/revisions/cumulative").json()
    assert len(rows) == len(revised_rows())
    by_tag = {r["tag"]: r for r in rows}
    assert (
        by_tag["320-DI-000301"]["first_label"] == "01" and by_tag["320-DI-000301"]["changes"] == 1
    )


def test_labels_default_to_the_next_number_and_cannot_repeat(client, tmp_path):
    pid, _ = baselined(client, tmp_path)
    revise(client, tmp_path, pid, revised_rows())
    revise(client, tmp_path, pid, revised_rows() + cable_rows("320-CE-0004", 3))
    assert [e["label"] for e in client.get(f"/api/projects/{pid}/revisions").json()] == [
        "00", "01", "02",
    ]  # fmt: skip
    dup = client.post(f"/api/projects/{pid}/revisions", json={"label": "1"})
    assert dup.status_code == 409 and dup.json()["code"] == "revision_exists"
    named = revise(client, tmp_path, pid, revised_rows(), label="05a")
    assert named["summary"]["revision"] == "05A"


def test_revision_without_uploads_is_refused(client, tmp_path):
    pid, _ = baselined(client, tmp_path)
    for u in client.get(f"/api/projects/{pid}/uploads").json():
        client.delete(f"/api/projects/{pid}/uploads/{u['id']}")
    r = client.post(f"/api/projects/{pid}/revisions", json={})
    assert r.status_code == 409 and r.json()["code"] == "no_uploads"


def test_revision_workbook_carries_the_impact_and_diff_sheets(client, tmp_path):
    sample_template_or_skip()
    from io import BytesIO

    from openpyxl import load_workbook

    pid, _ = baselined(client, tmp_path)
    done = revise(client, tmp_path, pid, revised_rows())
    r = client.get(f"/api/projects/{pid}/runs/{done['id']}/download")
    assert r.status_code == 200
    names = load_workbook(BytesIO(r.content), read_only=True).sheetnames
    assert "Revision Impact" in names and "Revision Diff" in names


def test_baseline_from_an_issued_workbook(client, tmp_path, sample_alloc_path):
    pid = new_project(client)
    ok = client.post(
        f"/api/projects/{pid}/baseline/upload",
        files={"file": ("issued.xlsx", sample_alloc_path.read_bytes(), "application/octet-stream")},
    )
    assert ok.status_code == 201, ok.text
    assert ok.json()["label"] == "00" and ok.json()["placed"] > 0
    assert [e["label"] for e in client.get(f"/api/projects/{pid}/revisions").json()] == ["00"]
    again = client.post(
        f"/api/projects/{pid}/baseline/upload",
        files={"file": ("issued.xlsx", sample_alloc_path.read_bytes(), "application/octet-stream")},
    )
    assert again.status_code == 409 and again.json()["code"] == "baseline_exists"


def test_a_workbook_that_is_not_an_allocation_is_a_422(client, tmp_path):
    pid = new_project(client)
    bad = client.post(
        f"/api/projects/{pid}/baseline/upload",
        files={
            "file": ("mdb.xlsx", xlsx_bytes(tmp_path, simple_rows()), "application/octet-stream")
        },
    )
    assert bad.status_code == 422 and bad.json()["code"] == "not_an_allocation"
    wrong = client.post(
        f"/api/projects/{pid}/baseline/upload", files={"file": ("x.txt", b"hi", "text/plain")}
    )
    assert wrong.status_code == 415


def test_uploads_can_be_removed(client, tmp_path):
    pid = new_project(client)
    up = upload(client, pid, xlsx_bytes(tmp_path, simple_rows())).json()
    assert client.delete(f"/api/projects/{pid}/uploads/{up['id']}").status_code == 204
    assert client.get(f"/api/projects/{pid}/uploads").json() == []
    gone = client.delete(f"/api/projects/{pid}/uploads/{up['id']}")
    assert gone.status_code == 404 and gone.json()["code"] == "not_found"


def test_other_projects_chains_are_separate(client, tmp_path):
    pid, _ = baselined(client, tmp_path)
    other = new_project(client)
    assert client.get(f"/api/projects/{other}/revisions").json() == []
    assert client.get(f"/api/projects/{other}/revisions/00").status_code == 404
    assert pid != other


def test_unexpected_fields_are_ignored_and_unknown_paths_have_the_error_shape(client):  # T13.1.3/4
    pid = new_project(client, extra_field="ignored")
    r = client.post(f"/api/projects/{pid}/revisions", json={"label": "01", "surprise": True})
    assert r.status_code == 409  # got past validation: the extra key was ignored
    nope = client.get("/api/no/such/path")
    assert nope.status_code == 404
    assert set(nope.json()) == {"code", "message", "details"}


def test_cors_is_open_to_local_origins_by_default_and_closed_in_release(tmp_path):
    from app.config import Settings
    from tests.api_helpers import TestClient, create_app

    def preflight(settings, origin):
        c = TestClient(create_app(settings))
        return c.options(
            "/api/projects",
            headers={"Origin": origin, "Access-Control-Request-Method": "POST"},
        ).headers.get("access-control-allow-origin")

    dev = Settings(db_path=str(tmp_path / "a.db"), cpsat_enabled=False)
    assert preflight(dev, "http://localhost:5173") == "http://localhost:5173"
    assert preflight(dev, "https://evil.example") is None
    rel = Settings(
        db_path=str(tmp_path / "b.db"), cpsat_enabled=False, cors_origins=("https://app.example",)
    )
    assert preflight(rel, "https://app.example") == "https://app.example"
    assert preflight(rel, "http://localhost:5173") is None


def test_cors_origins_setting_reads_a_comma_list(monkeypatch):
    from app.config import Settings

    monkeypatch.setenv("CORS_ORIGINS", "https://a.example/, https://b.example")
    assert Settings.from_env().cors_origins == ("https://a.example", "https://b.example")
    monkeypatch.delenv("CORS_ORIGINS")
    assert Settings.from_env().cors_origins == ()
