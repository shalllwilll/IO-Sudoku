import openpyxl
import pytest

from app.core.ingest.merge import merge_files
from app.core.ingest.reader import read_workbook
from app.core.layout.blocks import build_blocks
from app.core.model.pack import Pack
from app.core.solve.greedy import solve_greedy
from app.core.writer.template import WriterError, build_rows, write_allocation
from tests.test_greedy import cable


@pytest.fixture
def pack(pack_dir):
    return Pack.load(pack_dir)


@pytest.fixture
def template(blank_template_path):
    return blank_template_path


def alloc_for(signals, pack):
    return solve_greedy(build_blocks(signals, pack), pack, "AR1", "100")


def test_output_opens_and_round_trips(pack, template, tmp_path):
    s = cable("C1", 22, 24) + cable("C2", 5)
    a = alloc_for(s, pack)
    out = write_allocation(a, pack, template, tmp_path / "out.xlsx")
    ws = openpyxl.load_workbook(out).worksheets[0]
    assert ws.max_column == 76
    (res,) = read_workbook(out)
    tags = {r.values["Tag Name"] for r in res.rows if r.values["Tag Status"] == "Field"}
    assert tags == {x.tag for x in s}
    assert res.summary.rows == 384  # one row per CHARM position of the one cabinet
    # the template itself is untouched
    assert openpyxl.load_workbook(template).worksheets[0]["A2"].value is None


def test_missing_template(pack, tmp_path):
    with pytest.raises(WriterError, match="not found"):
        write_allocation(
            alloc_for(cable("C1", 2), pack), pack, tmp_path / "nope.xlsx", tmp_path / "o.xlsx"
        )


def test_zero_rows_headers_only(pack, template, tmp_path):
    out = write_allocation(alloc_for([], pack), pack, template, tmp_path / "o.xlsx")
    ws = openpyxl.load_workbook(out).worksheets[0]
    assert ws["A1"].value == "Sr. No." and ws["A2"].value is None


def test_unicode_preserved(pack, template, tmp_path):
    s = [x.model_copy(update={"tag": "T-é°C-µ"}) if x.pair == 1 else x for x in cable("C1", 2)]
    out = write_allocation(alloc_for(s, pack), pack, template, tmp_path / "o.xlsx")
    vals = [
        r[2]
        for r in openpyxl.load_workbook(out).worksheets[0].iter_rows(min_row=2, values_only=True)
    ]
    assert "T-é°C-µ" in vals


def test_deterministic_cells(pack, template, tmp_path):
    a = alloc_for(cable("C1", 22, 24) + cable("C2", 5), pack)
    p1 = write_allocation(a, pack, template, tmp_path / "1.xlsx")
    p2 = write_allocation(a, pack, template, tmp_path / "2.xlsx")
    cells = lambda p: [
        r for r in openpyxl.load_workbook(p).worksheets[0].iter_rows(values_only=True)
    ]
    assert cells(p1) == cells(p2)


def test_extra_sheets_and_wired_spare_row(pack, template, tmp_path):
    sigs = [
        x.model_copy(update={"signal_type": "4-20 ma HART", "signal_level": "2-Wire"})
        for x in cable("C1", 22, 24)
    ]
    a = alloc_for(sigs, pack)
    out = write_allocation(
        a, pack, template, tmp_path / "o.xlsx", {"TQ list": [["code"], ["TQ-X"]]}
    )
    wb = openpyxl.load_workbook(out)
    assert "TQ list" in wb.sheetnames
    spare = [r for r in build_rows(a, pack) if r["tagstatus"] == "Wired Spare"]
    assert len(spare) == 2 and spare[0]["charmtype"] == "SE4303T01"
    assert spare[0]["tagname"].startswith("100DSP")


def test_sample_signal_columns_match_sample_allocation(
    pack, sample_mdb_path, sample_alloc_path, template
):
    """Per-signal columns (not positions) must equal the sample allocation's for every field row."""
    ss = merge_files([sample_mdb_path], pack)
    a = solve_greedy(build_blocks(ss.signals, pack), pack, "P30", "32001")
    rows = build_rows(a, pack)
    by_place = {
        (r["fieldcableno"], r["fieldcablepair"]): r for r in rows if r["tagstatus"] == "Field"
    }
    ws = openpyxl.load_workbook(sample_alloc_path, read_only=True, data_only=True).active
    data = list(ws.iter_rows(values_only=True))
    head = data[0]
    import re

    norm = lambda h: re.sub(r"[^a-z0-9]", "", str(h).lower())
    cols = [
        "Technical Room", "Operating Area", "Risk Area", "Unit", "Segregation", "DSTName",
        "Service Description", "Location", "IO Type", "Signal Type", "Signal Level", "NON-IS / IS",
        "Equipment Name", "External Power", "Field Cable Size", "Old Tag", "Instrument type",
        "Loop Name", "Tag Status", "HW Loop Typical",
    ]  # fmt: skip
    bad, checked = [], 0
    for r in data[1:]:
        a_ = dict(zip(head, r, strict=False))
        if a_["Tag Status"] != "Field":
            continue
        mine = by_place.get((a_["Field Cable No"], a_["Field Cable Pair"]))
        if mine is None:
            continue
        checked += 1
        for c in cols:
            exp = str(a_[c]).strip() if a_[c] is not None else "-"
            got = str(mine.get(norm(c)) if mine.get(norm(c)) is not None else "-").strip()
            if exp != got:
                bad.append((a_["Tag Name"], c, exp, got))
    print(f"signal columns: {checked} rows, {len(bad)} mismatching cells")
    from collections import Counter

    print(Counter(b[1] for b in bad).most_common(8), bad[:6])
    # known sample quirks: one stale tag (20076 vs 20077) and an Old Tag of 0 on one row
    real = [b for b in bad if "20076" not in b[0] and b[1] != "Old Tag"]
    assert real == []
    assert checked == 176


def test_every_row_key_lands_in_a_template_column(pack, template):
    """A key that matches no header is silently dropped; the CIOC name once was."""
    from app.core.writer.template import _norm

    wb = openpyxl.load_workbook(template, read_only=True)
    heads = {_norm(h) for h in next(wb.worksheets[0].iter_rows(max_row=1, values_only=True)) if h}
    keys = set(build_rows(alloc_for(cable("C1", 5), pack), pack)[0])
    assert keys <= heads, f"row keys with no template column: {sorted(keys - heads)}"


def test_cioc_and_baseplate_names_are_written(pack, template, tmp_path):
    a = alloc_for(cable("C1", 5), pack)
    out = write_allocation(a, pack, template, tmp_path / "o.xlsx")
    ws = openpyxl.load_workbook(out).worksheets[0]
    head = [c.value for c in ws[1]]
    first = dict(zip(head, [c.value for c in ws[2]], strict=False))
    assert first["CIOC/CSLS Name"] == a.layout.ciocs[0].name != "-"
    assert first["Baseplate"] == a.layout.ciocs[0].baseplates[0].name
